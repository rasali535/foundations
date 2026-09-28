from io import BytesIO

import pytest
from fastapi import UploadFile
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen import canvas

from routers.admin_router import _parse_roster_upload


@pytest.mark.asyncio
async def test_parse_excel_roster_first_surname_email():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["First name", "Surname", "Email address"])
    sheet.append(["Jane", "Doe", "JANE@EXAMPLE.COM"])
    sheet.append(["John", "Smith", "john@example.com"])

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    contacts = await _parse_roster_upload(
        UploadFile(filename="employees.xlsx", file=buffer)
    )

    assert [(c.name, c.email) for c in contacts] == [
        ("Jane Doe", "jane@example.com"),
        ("John Smith", "john@example.com"),
    ]


@pytest.mark.asyncio
async def test_parse_docx_roster_table_first_surname_email():
    document = Document()
    table = document.add_table(rows=1, cols=3)
    headers = table.rows[0].cells
    headers[0].text = "First name"
    headers[1].text = "Surname"
    headers[2].text = "Email address"

    row = table.add_row().cells
    row[0].text = "Naledi"
    row[1].text = "Molefe"
    row[2].text = "naledi@example.com"

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    contacts = await _parse_roster_upload(
        UploadFile(filename="employees.docx", file=buffer)
    )

    assert len(contacts) == 1
    assert contacts[0].name == "Naledi Molefe"
    assert contacts[0].email == "naledi@example.com"


@pytest.mark.asyncio
async def test_parse_pdf_roster_lines_first_surname_email():
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 760, "First name, Surname, Email address")
    pdf.drawString(72, 740, "Kagiso, Dube, kagiso@example.com")
    pdf.drawString(72, 720, "Thato, Moyo, thato@example.com")
    pdf.save()
    buffer.seek(0)

    contacts = await _parse_roster_upload(
        UploadFile(filename="employees.pdf", file=buffer)
    )

    assert {(c.name, c.email) for c in contacts} == {
        ("Kagiso Dube", "kagiso@example.com"),
        ("Thato Moyo", "thato@example.com"),
    }


@pytest.mark.asyncio
async def test_parse_roster_deduplicates_email():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["First name", "Surname", "Email address"])
    sheet.append(["Jane", "Doe", "jane@example.com"])
    sheet.append(["Janet", "Doe", "JANE@EXAMPLE.COM"])

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    contacts = await _parse_roster_upload(
        UploadFile(filename="employees.xlsx", file=buffer)
    )

    assert len(contacts) == 1
    assert contacts[0].email == "jane@example.com"
